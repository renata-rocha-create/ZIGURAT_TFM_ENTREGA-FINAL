"""
montar_planilha.py — Monta benchmark_resultados.xlsx a partir dos CSVs das rodadas.

Entradas (todas em benchmark/):
  resultados.csv, erros_auditor.csv          → rodada v2 (referência)
  rodada_v1/resultados.csv                    → rodada v1 (antes da correção da bacia)
  gabarito.csv, gabarito_errata.csv

As métricas da planilha são FÓRMULAS do Excel (CONT.SES sobre as abas de
classificação), para que a banca possa auditar cada número. Depois de gerar,
abra no Excel (ou rode o recalc) para ver os valores.

Uso: python benchmark/montar_planilha.py
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

for _f in (sys.stdout, sys.stderr):
    try:
        _f.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

AQUI = Path(__file__).resolve().parent
ITENS = ["4.6.6", "5.4.3", "6.3.4", "6.6", "6.11.1", "6.11.2", "6.11.3",
         "7.5", "7.6-7.8", "7.7.1", "7.7.2.1", "7.8"]
CAMADAS = ["python", "llm", "final"]

FONTE = "Arial"
F_BASE = Font(name=FONTE, size=10)
F_NEG = Font(name=FONTE, size=10, bold=True)
F_TIT = Font(name=FONTE, size=14, bold=True)
F_SUB = Font(name=FONTE, size=11, bold=True)
F_CAB = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
F_NOTA = Font(name=FONTE, size=9, italic=True, color="595959")
P_CAB = PatternFill("solid", start_color="1F3864")
P_SUB = PatternFill("solid", start_color="D9E1F2")
P_DEST = PatternFill("solid", start_color="FFF2CC")
FINO = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
QUEBRA = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(horizontal="center", vertical="center")


def ler(caminho):
    with open(caminho, encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def cabecalho(ws, linha, titulos, col0=1):
    for i, t in enumerate(titulos):
        c = ws.cell(row=linha, column=col0 + i, value=t)
        c.font, c.fill, c.border = F_CAB, P_CAB, BORDA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def celula(ws, linha, col, valor, fonte=F_BASE, fmt=None, borda=True, alin=None):
    c = ws.cell(row=linha, column=col, value=valor)
    c.font = fonte
    if borda:
        c.border = BORDA
    if fmt:
        c.number_format = fmt
    if alin:
        c.alignment = alin
    return c


def larguras(ws, mapa):
    for col, w in mapa.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------- abas de classificação
def aba_classificacao(wb, nome, linhas, errata):
    """Uma linha por par variante × item, com a classe (VP/FP/FN/VN) calculada por fórmula."""
    ws = wb.create_sheet(nome)
    tit = ["Variante", "Item NBR", "Esperado (gabarito)", "Esperado (errata)",
           "Python", "LLM", "Final", "Sensib. W01 (1 = excluir)",
           "Python estrito", "Python amplo", "LLM estrito", "LLM amplo",
           "Final estrito", "Final amplo",
           "Python exato", "LLM exato", "Final exato"]
    cabecalho(ws, 1, tit)
    pos_e = 'OR({x}="Não Conforme",{x}="Parcial")'
    pos_a = 'OR({x}="Não Conforme",{x}="Parcial",{x}="Indeterminado")'
    real = 'AND($C{r}<>"Conforme",$C{r}<>"N/A")'
    for i, d in enumerate(linhas, start=2):
        esp_ag = errata.get((d["variante"], d["item_nbr"]), d["status_esperado"])
        valores = [d["variante"], d["item_nbr"], d["status_esperado"], esp_ag,
                   d["python"], d["llm"], d["final"],
                   1 if (d["variante"], d["item_nbr"]) == ("W01", "7.7.1") else 0]
        for j, v in enumerate(valores, start=1):
            celula(ws, i, j, v)
        if esp_ag != d["status_esperado"]:
            ws.cell(row=i, column=4).fill = P_DEST
        for k, col_cam in enumerate("EFG"):
            x = f"{col_cam}{i}"
            r = real.format(r=i)
            f_est = (f'=IF({x}="—","—",IF($C{i}="Indeterminado","fora",'
                     f'IF({r},IF({pos_e.format(x=x)},"VP","FN"),IF({pos_e.format(x=x)},"FP","VN"))))')
            f_amp = (f'=IF({x}="—","—",'
                     f'IF({r},IF({pos_a.format(x=x)},"VP","FN"),IF({pos_a.format(x=x)},"FP","VN")))')
            f_exa = f'=IF({x}="—","—",IF({x}=$D{i},1,0))'
            celula(ws, i, 9 + 2 * k, f_est, alin=CENTRO)
            celula(ws, i, 10 + 2 * k, f_amp, alin=CENTRO)
            celula(ws, i, 15 + k, f_exa, alin=CENTRO)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:Q{len(linhas) + 1}"
    larguras(ws, {"A": 9, "B": 9, "C": 15, "D": 15, "E": 14, "F": 14, "G": 14, "H": 11,
                  **{get_column_letter(c): 10 for c in range(9, 18)}})
    return len(linhas) + 1


# colunas de classe por camada/critério nas abas de classificação
COL = {("python", "estrito"): "I", ("python", "amplo"): "J",
       ("llm", "estrito"): "K", ("llm", "amplo"): "L",
       ("final", "estrito"): "M", ("final", "amplo"): "N"}
COL_EXATO = {"python": "O", "llm": "P", "final": "Q"}


def bloco_metricas(ws, linha, col0, aba, ult, camada, crit, filtro=""):
    """Escreve n, VP, FP, FN, VN, acurácia, precisão, recall, F1 a partir de CONT.SES."""
    rng = f"'{aba}'!${COL[(camada, crit)]}$2:${COL[(camada, crit)]}${ult}"
    refs = []
    for k, cls in enumerate(["VP", "FP", "FN", "VN"]):
        c = get_column_letter(col0 + 1 + k)
        refs.append(f"{c}{linha}")
        celula(ws, linha, col0 + 1 + k, f'=COUNTIFS({rng},"{cls}"{filtro})', alin=CENTRO)
    vp, fp, fn, vn = refs
    celula(ws, linha, col0, f"={vp}+{fp}+{fn}+{vn}", alin=CENTRO)
    n = f"{get_column_letter(col0)}{linha}"
    p = f"{get_column_letter(col0 + 6)}{linha}"
    r = f"{get_column_letter(col0 + 7)}{linha}"
    celula(ws, linha, col0 + 5, f'=IFERROR(({vp}+{vn})/{n},"—")', fmt="0.000", alin=CENTRO)
    celula(ws, linha, col0 + 6, f'=IFERROR({vp}/({vp}+{fp}),"—")', fmt="0.000", alin=CENTRO)
    celula(ws, linha, col0 + 7, f'=IFERROR({vp}/({vp}+{fn}),"—")', fmt="0.000", alin=CENTRO)
    celula(ws, linha, col0 + 8, f'=IFERROR(2*{p}*{r}/({p}+{r}),"—")', fmt="0.000", alin=CENTRO)


METR = ["n", "VP", "FP", "FN", "VN", "Acurácia", "Precisão", "Recall", "F1"]


def aba_resumo(wb, ult2, ult1):
    ws = wb.create_sheet("Resumo", 1)
    celula(ws, 1, 1, "Benchmark do Auditor NBR 9050 — resultados (30 variantes × 12 itens = 360 pares)",
           F_TIT, borda=False)
    celula(ws, 2, 1, "Todas as métricas abaixo são fórmulas sobre as abas 'Classificação v2' e "
           "'Classificação v1'. v2 = rodada de referência (após a correção da altura da bacia).",
           F_NOTA, borda=False)

    # 1) v2 por camada
    celula(ws, 4, 1, "1. Rodada v2 (referência) — métricas globais por camada", F_SUB, borda=False)
    cabecalho(ws, 5, ["Camada", "Critério"] + METR)
    lin = 6
    for cam in CAMADAS:
        for crit in ("estrito", "amplo"):
            celula(ws, lin, 1, cam, F_NEG)
            celula(ws, lin, 2, crit)
            bloco_metricas(ws, lin, 3, "Classificação v2", ult2, cam, crit)
            if cam == "final":
                for c in range(1, 12):
                    ws.cell(row=lin, column=c).fill = P_DEST
            lin += 1
    celula(ws, lin, 1, "Python não mede o 7.7.1 (30 pares '—'); por isso n é menor nessa camada. "
           "No critério estrito, pares com esperado 'Indeterminado' ficam fora (2 pares).",
           F_NOTA, borda=False)

    # 2) acurácia de status exato
    lin += 2
    celula(ws, lin, 1, "2. Acurácia de status exato (acertou a categoria; gabarito com errata)", F_SUB, borda=False)
    lin += 1
    cabecalho(ws, lin, ["Camada", "Pares avaliados", "Acertos v2", "Exato v2", "Acertos v1", "Exato v1"])
    for cam in CAMADAS:
        lin += 1
        c = COL_EXATO[cam]
        celula(ws, lin, 1, cam, F_NEG)
        celula(ws, lin, 2, f"=COUNT('Classificação v2'!{c}2:{c}{ult2})", alin=CENTRO)
        celula(ws, lin, 3, f"=SUM('Classificação v2'!{c}2:{c}{ult2})", alin=CENTRO)
        celula(ws, lin, 4, f"=C{lin}/B{lin}", fmt="0.000", alin=CENTRO)
        celula(ws, lin, 5, f"=SUM('Classificação v1'!{c}2:{c}{ult1})", alin=CENTRO)
        celula(ws, lin, 6, f"=E{lin}/COUNT('Classificação v1'!{c}2:{c}{ult1})", fmt="0.000", alin=CENTRO)

    # 3) v1 × v2 lado a lado
    lin += 3
    celula(ws, lin, 1, "3. Rodada v1 × v2 lado a lado (critérios estrito e amplo)", F_SUB, borda=False)
    lin += 1
    celula(ws, lin, 3, "v1 (antes da correção)", F_NEG, borda=False)
    celula(ws, lin, 12, "v2 (referência)", F_NEG, borda=False)
    lin += 1
    cabecalho(ws, lin, ["Camada", "Critério"] + METR + METR)
    for cam in CAMADAS:
        for crit in ("estrito", "amplo"):
            lin += 1
            celula(ws, lin, 1, cam, F_NEG)
            celula(ws, lin, 2, crit)
            bloco_metricas(ws, lin, 3, "Classificação v1", ult1, cam, crit)
            bloco_metricas(ws, lin, 12, "Classificação v2", ult2, cam, crit)
    lin += 1
    celula(ws, lin, 1, "Python v1→v2: só mudou S02/7.7.2.1 (Indeterminado → Parcial), efeito da "
           "correção da altura da bacia. Diferenças na camada LLM vêm do não determinismo do modelo.",
           F_NOTA, borda=False)

    # 4) sensibilidade W01
    lin += 3
    celula(ws, lin, 1, "4. Análise de sensibilidade — camada final v2 sem o par W01 / 7.7.1", F_SUB, borda=False)
    lin += 1
    celula(ws, lin, 1, "A mutação W01 só alterou o contorno do ambiente (giro ⌀1,50 m). O LLM acusou "
           "7.7.1 (área de transferência), que o Python não mede — o 'erro' é ambíguo: o espaço "
           "realmente ficou menor.", F_NOTA, borda=False)
    lin += 1
    cabecalho(ws, lin, ["Camada", "Critério"] + METR)
    for crit in ("estrito", "amplo"):
        lin += 1
        celula(ws, lin, 1, "final sem W01/7.7.1", F_NEG)
        celula(ws, lin, 2, crit)
        bloco_metricas(ws, lin, 3, "Classificação v2", ult2, "final", crit,
                       filtro=f",'Classificação v2'!$H$2:$H${ult2},0")

    # 5) ganho da arquitetura híbrida
    lin += 3
    celula(ws, lin, 1, "5. O que a arquitetura híbrida ('LLM traduz, Python julga') ganha — v2, critério estrito",
           F_SUB, borda=False)
    lin += 1
    cabecalho(ws, lin, ["Métrica", "Só LLM", "Final (híbrido)", "Diferença"])
    base = 6  # linhas do bloco 1: python estrito=6, amplo=7, llm estrito=8, amplo=9, final estrito=10, amplo=11
    for nome, col in (("Precisão", "I"), ("Recall", "J"), ("F1", "K"), ("Alarmes falsos (FP)", "E")):
        lin += 1
        celula(ws, lin, 1, nome, F_NEG)
        fmt = "0" if nome.startswith("Alarmes") else "0.000"
        celula(ws, lin, 2, f"={col}{base + 2}", fmt=fmt, alin=CENTRO)
        celula(ws, lin, 3, f"={col}{base + 4}", fmt=fmt, alin=CENTRO)
        celula(ws, lin, 4, f"=C{lin}-B{lin}", fmt=("+0;-0;0" if fmt == "0" else "+0.000;-0.000;0.000"),
               alin=CENTRO)

    larguras(ws, {"A": 20, "B": 11, **{get_column_letter(c): 9 for c in range(3, 21)}})
    ws.freeze_panes = "A4"


def aba_por_item(wb, ult2):
    ws = wb.create_sheet("Métricas por item")
    celula(ws, 1, 1, "Rodada v2 — métricas por item NBR (fórmulas sobre 'Classificação v2')", F_TIT, borda=False)
    lin = 3
    for crit in ("estrito", "amplo"):
        celula(ws, lin, 1, f"Critério {crit}", F_SUB, borda=False)
        lin += 1
        cabecalho(ws, lin, ["Item NBR", "Camada"] + METR)
        for item in ITENS:
            for cam in CAMADAS:
                lin += 1
                celula(ws, lin, 1, item, F_NEG)
                celula(ws, lin, 2, cam)
                bloco_metricas(ws, lin, 3, "Classificação v2", ult2, cam, crit,
                               filtro=f",'Classificação v2'!$B$2:$B${ult2},\"{item}\"")
                if cam == "final":
                    for c in range(1, 12):
                        ws.cell(row=lin, column=c).fill = P_SUB
        lin += 3
    celula(ws, lin - 1, 1, "n = 0 indica que a camada não avalia o item (Python no 7.7.1). "
           "'—' em precisão/recall: divisão por zero (sem positivos previstos ou reais).",
           F_NOTA, borda=False)
    larguras(ws, {"A": 10, "B": 9, **{get_column_letter(c): 9 for c in range(3, 12)}})
    ws.freeze_panes = "C3"


def aba_alarmes_llm(wb, ult1, ult2):
    ws = wb.create_sheet("Alarmes falsos LLM")
    celula(ws, 1, 1, "Onde o LLM sozinho erra — falsos positivos (FP) e negativos (FN) por item", F_TIT, borda=False)
    celula(ws, 2, 1, "Contagens por CONT.SES. Em todos esses pares a camada final (Python) corrige o LLM, "
           "exceto W01/7.7.1 (item que o Python não mede).", F_NOTA, borda=False)
    cabecalho(ws, 4, ["Item NBR", "FP estrito v1", "FP estrito v2", "FP amplo v1", "FP amplo v2",
                      "FN estrito v2", "FP final v2 (amplo)", "Causa provável (LLM)"])
    causas = {
        "7.7.2.1": "Sem altura explícita da bacia no texto, o LLM marca Indeterminado/Parcial; o Python lê a geometria.",
        "4.6.6": "Maçaneta tipo alavanca está nos Psets do TIPO da porta, que o prompt não inclui → Indeterminado.",
        "7.6-7.8": "Regras de barra por tipo (A, A1, lavatório) exigem cotas 3D; o LLM tende a Parcial.",
        "6.3.4": "Desnível de 5 mm no piso: o LLM hesita entre tolerância e falha → Indeterminado.",
        "6.6": "Inclinação calculada pelo Python (laje inclinada); o LLM às vezes não reconhece a rampa.",
        "6.11.3": "Caso isolado (variação de uma rodada).",
        "7.7.1": "W01: o LLM associou o ambiente reduzido à área de transferência (ambíguo).",
    }
    lin = 4
    for item in ITENS:
        lin += 1
        f_it = lambda aba, u: f",'{aba}'!$B$2:$B${u},\"{item}\""
        celula(ws, lin, 1, item, F_NEG)
        celula(ws, lin, 2, f"=COUNTIFS('Classificação v1'!$K$2:$K${ult1},\"FP\"{f_it('Classificação v1', ult1)})", alin=CENTRO)
        celula(ws, lin, 3, f"=COUNTIFS('Classificação v2'!$K$2:$K${ult2},\"FP\"{f_it('Classificação v2', ult2)})", alin=CENTRO)
        celula(ws, lin, 4, f"=COUNTIFS('Classificação v1'!$L$2:$L${ult1},\"FP\"{f_it('Classificação v1', ult1)})", alin=CENTRO)
        celula(ws, lin, 5, f"=COUNTIFS('Classificação v2'!$L$2:$L${ult2},\"FP\"{f_it('Classificação v2', ult2)})", alin=CENTRO)
        celula(ws, lin, 6, f"=COUNTIFS('Classificação v2'!$K$2:$K${ult2},\"FN\"{f_it('Classificação v2', ult2)})", alin=CENTRO)
        celula(ws, lin, 7, f"=COUNTIFS('Classificação v2'!$N$2:$N${ult2},\"FP\"{f_it('Classificação v2', ult2)})", alin=CENTRO)
        celula(ws, lin, 8, causas.get(item, ""), alin=QUEBRA)
    lin += 1
    celula(ws, lin, 1, "Total", F_NEG)
    for c in range(2, 8):
        L = get_column_letter(c)
        celula(ws, lin, c, f"=SUM({L}5:{L}{lin - 1})", F_NEG, alin=CENTRO)
    lin += 2
    celula(ws, lin, 1, "Leitura: no critério amplo, 'Indeterminado' conta como alarme. A maioria dos FP do LLM "
           "é Indeterminado — ele pede revisão humana quando o dado não está explícito no texto.",
           F_NOTA, borda=False)
    larguras(ws, {"A": 10, "B": 11, "C": 11, "D": 11, "E": 11, "F": 11, "G": 13, "H": 80})


def aba_consistencia(wb, res1, res2):
    ws = wb.create_sheet("Consistência LLM v1×v2")
    celula(ws, 1, 1, "Repetibilidade do LLM: mesmo IFC, mesmo prompt, temperatura 0, duas rodadas", F_TIT, borda=False)
    v1 = {(d["variante"], d["item_nbr"]): d for d in res1}
    # resumo no topo
    n = len(res2)
    ini, fim = 12, 12 + n - 1
    celula(ws, 3, 1, "Pares comparados", F_NEG)
    celula(ws, 3, 2, f"=COUNTA(A{ini}:A{fim})", alin=CENTRO)
    celula(ws, 4, 1, "LLM igual nas 2 rodadas", F_NEG)
    celula(ws, 4, 2, f"=SUM(E{ini}:E{fim})", alin=CENTRO)
    celula(ws, 5, 1, "Consistência LLM", F_NEG)
    celula(ws, 5, 2, "=B4/B3", fmt="0.0%", alin=CENTRO).fill = P_DEST
    celula(ws, 6, 1, "Python igual nas 2 rodadas", F_NEG)
    celula(ws, 6, 2, f"=SUM(H{ini}:H{fim})", alin=CENTRO)
    celula(ws, 7, 1, "Consistência Python", F_NEG)
    celula(ws, 7, 2, "=B6/B3", fmt="0.0%", alin=CENTRO)
    celula(ws, 4 + len(ITENS) + 1, 1, "O Python só mudou em S02/7.7.2.1 (correção do código entre as rodadas, "
           "não variação aleatória).", F_NOTA, borda=False)
    # divergências por item
    cabecalho(ws, 3, ["Item NBR", "Divergências LLM"], col0=4)
    for k, item in enumerate(ITENS):
        celula(ws, 4 + k, 4, item, F_NEG)
        celula(ws, 4 + k, 5, f'=COUNTIFS(B{ini}:B{fim},"{item}",E{ini}:E{fim},0)', alin=CENTRO)
    # cuidado: tabela de itens ocupa linhas 4..15 nas colunas D:E → dados começam após ela
    ini = 4 + len(ITENS) + 3
    fim = ini + n - 1
    # reescrever as fórmulas do resumo com o intervalo certo
    ws["B3"] = f"=COUNTA(A{ini}:A{fim})"
    ws["B4"] = f"=SUM(E{ini}:E{fim})"
    ws["B6"] = f"=SUM(H{ini}:H{fim})"
    for k, item in enumerate(ITENS):
        ws.cell(row=4 + k, column=5).value = f'=COUNTIFS(B{ini}:B{fim},"{item}",E{ini}:E{fim},0)'
    cabecalho(ws, ini - 1, ["Variante", "Item NBR", "LLM v1", "LLM v2", "LLM igual",
                            "Python v1", "Python v2", "Python igual"])
    for i, d in enumerate(res2):
        r = ini + i
        a = v1[(d["variante"], d["item_nbr"])]
        for j, v in enumerate([d["variante"], d["item_nbr"], a["llm"], d["llm"]], start=1):
            celula(ws, r, j, v)
        celula(ws, r, 5, f"=IF(C{r}=D{r},1,0)", alin=CENTRO)
        celula(ws, r, 6, a["python"])
        celula(ws, r, 7, d["python"])
        celula(ws, r, 8, f"=IF(F{r}=G{r},1,0)", alin=CENTRO)
    ws.auto_filter.ref = f"A{ini - 1}:H{fim}"
    larguras(ws, {"A": 26, "B": 10, "C": 14, "D": 14, "E": 16, "F": 14, "G": 14, "H": 12})


def aba_erros(wb, erros):
    ws = wb.create_sheet("Erros do auditor v2")
    celula(ws, 1, 1, "Divergências da rodada v2 contra o gabarito com errata", F_TIT, borda=False)
    celula(ws, 2, 1, "Camadas python e final no topo (com causa); camada llm abaixo. Fonte: erros_auditor.csv.",
           F_NOTA, borda=False)
    causas = {
        ("S01", "7.7.2.1"): "Regra de projeto: bacia fora de 0,41–0,47 m vira 'verificar in loco' "
                            "(Indeterminado), nunca Não Conforme. FN estrito, acerto no amplo.",
        ("S02", "7.6-7.8"): "Barra de fundo a 3,2 cm da tampa (mín. 4 cm) aceita pela tolerância de "
                            "modelagem de 1 cm (TOL_BARRA_CONF). FN nos dois critérios.",
        ("W01", "7.7.1"): "Python não mede 7.7.1; vale o LLM, que ligou o ambiente reduzido à área de "
                          "transferência. Ambíguo (ver sensibilidade no Resumo).",
    }
    cabecalho(ws, 4, ["Camada", "Variante", "Item NBR", "Esperado", "Obtido", "Tipo de erro",
                      "Causa", "Detalhe do Python"])
    ordem = sorted(erros, key=lambda e: ({"python": 0, "final": 1, "llm": 2}[e["camada"]], e["variante"], e["item_nbr"]))
    for i, e in enumerate(ordem, start=5):
        causa = causas.get((e["variante"], e["item_nbr"]), "") if e["camada"] != "llm" else ""
        for j, v in enumerate([e["camada"], e["variante"], e["item_nbr"], e["esperado"], e["obtido"],
                               e["tipo_erro"], causa, e["detalhe_python"]], start=1):
            celula(ws, i, j, v, alin=QUEBRA)
        if e["camada"] != "llm":
            for j in range(1, 9):
                ws.cell(row=i, column=j).fill = P_DEST
    ws.auto_filter.ref = f"A4:H{len(ordem) + 4}"
    ws.freeze_panes = "A5"
    larguras(ws, {"A": 8, "B": 9, "C": 9, "D": 14, "E": 14, "F": 26, "G": 55, "H": 60})


def aba_tabela(wb, nome, linhas, titulo):
    ws = wb.create_sheet(nome)
    celula(ws, 1, 1, titulo, F_SUB, borda=False)
    campos = list(linhas[0].keys())
    cabecalho(ws, 3, campos)
    for i, d in enumerate(linhas, start=4):
        for j, k in enumerate(campos, start=1):
            celula(ws, i, j, d[k], alin=QUEBRA)
    ws.auto_filter.ref = f"A3:{get_column_letter(len(campos))}{len(linhas) + 3}"
    ws.freeze_panes = "A4"
    for j, k in enumerate(campos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = min(60, max(10, max(len(str(d[k])) for d in linhas[:60]) + 2))


def aba_leiame(wb):
    ws = wb.active
    ws.title = "LEIA-ME"
    textos = [
        ("Benchmark do Auditor BIM de Acessibilidade (NBR 9050:2020)", F_TIT),
        ("", None),
        ("O que é", F_SUB),
        ("Teste de injeção controlada de erros: 30 variantes IFC geradas por script a partir do modelo-base "
         "BENCHMARK_00_3.ifc (100% conforme), cada uma com erros conhecidos e tabelados ANTES da execução "
         "(gabarito pré-registrado, commit acaa304).", F_BASE),
        ("", None),
        ("Ambiente das rodadas (computador do projeto, 03/10/2026)", F_SUB),
        ("Windows · Python 3.12 · ifcopenshell 0.9.0 · LLM claude-sonnet-4-5, temperatura 0 · "
         "rodar_benchmark.py --llm anthropic", F_BASE),
        ("v1 = primeira rodada completa (revelou o bug da altura da bacia).  "
         "v2 = rodada após a correção — é a REFERÊNCIA da dissertação.", F_BASE),
        ("", None),
        ("Três camadas", F_SUB),
        ("python = verificação determinística (geometria e propriedades do IFC).", F_BASE),
        ("llm = veredito do modelo de linguagem sozinho.", F_BASE),
        ("final = o que o relatório mostra: o Python sobrescreve o LLM nos itens que mede "
         "('LLM traduz, Python julga'). O Python não mede o 7.7.1; ali vale o LLM.", F_BASE),
        ("", None),
        ("Definições (unidade = par variante × item; 30 × 12 = 360 pares)", F_SUB),
        ("Positivo real: status esperado diferente de Conforme e de N/A.", F_BASE),
        ("Critério ESTRITO: o auditor 'acusa' quando dá Não Conforme ou Parcial. Pares com esperado "
         "'Indeterminado' ficam fora.", F_BASE),
        ("Critério AMPLO: também conta Indeterminado ('pediu revisão humana').", F_BASE),
        ("Analogia do alarme de incêndio: VP = tocou com fogo · FP = alarme falso · FN = fogo sem alarme · "
         "VN = silêncio sem fogo.", F_BASE),
        ("Precisão = VP/(VP+FP): 'quando acusa, acerta?'.  Recall = VP/(VP+FN): 'dos problemas reais, "
         "quantos pegou?'.  F1 = média harmônica das duas.  Acurácia = (VP+VN)/n.", F_BASE),
        ("Acurácia de status exato: acertou a categoria (usa a errata por item, regra 'X de Y').", F_BASE),
        ("", None),
        ("Abas", F_SUB),
        ("Resumo — métricas globais v2, status exato, v1×v2 lado a lado, sensibilidade W01, ganho do híbrido.", F_BASE),
        ("Métricas por item — v2, por item NBR e camada, nos dois critérios.", F_BASE),
        ("Alarmes falsos LLM — FP/FN do LLM por item, v1 e v2, com a causa provável.", F_BASE),
        ("Consistência LLM v1×v2 — quantas respostas do LLM se repetiram entre as rodadas.", F_BASE),
        ("Erros do auditor v2 — cada divergência, com causa para python/final.", F_BASE),
        ("Classificação v2 / v1 — dados brutos + classe VP/FP/FN/VN calculada por FÓRMULA (base de tudo).", F_BASE),
        ("Gabarito / Errata — o gabarito pré-registrado e a errata gerada por regra (gerar_errata.py).", F_BASE),
        ("", None),
        ("Limitações", F_SUB),
        ("• Um único modelo-base, pequeno (3 ambientes, 31 pares positivos): as métricas valem para estes tipos de erro.", F_BASE),
        ("• Erros injetados por script não passam pela exportação do Revit (por isso o caso real C02 entra).", F_BASE),
        ("• W01: a mutação mudou só o contorno do ambiente; o FP do LLM no 7.7.1 é ambíguo (ver sensibilidade).", F_BASE),
        ("• O prompt do LLM não inclui os Psets do TIPO da porta — origem da maior parte dos 'Indeterminado' no 4.6.6.", F_BASE),
        ("• O LLM não é determinístico mesmo com temperatura 0 (ver consistência v1×v2).", F_BASE),
        ("• Os 2 FN do Python são decisões de tolerância (S01 bacia → Indeterminado; S02 tolerância de 1 cm da barra), "
         "não erros de medição.", F_BASE),
        ("", None),
        ("Gerada por benchmark/montar_planilha.py a partir dos CSVs do repositório.", F_NOTA),
    ]
    for i, (t, f) in enumerate(textos, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = f or F_BASE
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def main():
    res2 = ler(AQUI / "resultados.csv")
    res1 = ler(AQUI / "rodada_v1" / "resultados.csv")
    gab = ler(AQUI / "gabarito.csv")
    errata_l = ler(AQUI / "gabarito_errata.csv")
    erros = ler(AQUI / "erros_auditor.csv")
    errata = {(e["variante"], e["item_nbr"]): e["status_esperado_agregado"] for e in errata_l}

    wb = Workbook()
    aba_leiame(wb)
    ult2 = aba_classificacao(wb, "Classificação v2", res2, errata)
    ult1 = aba_classificacao(wb, "Classificação v1", res1, errata)
    aba_resumo(wb, ult2, ult1)
    aba_por_item(wb, ult2)
    aba_alarmes_llm(wb, ult1, ult2)
    aba_consistencia(wb, res1, res2)
    aba_erros(wb, erros)
    aba_tabela(wb, "Gabarito", gab, "Gabarito pré-registrado (gabarito.csv, commit acaa304) — não alterado")
    aba_tabela(wb, "Errata", errata_l, "Errata por regra (gerar_errata.py) — só afeta a acurácia de status exato")

    ordem = ["LEIA-ME", "Resumo", "Métricas por item", "Alarmes falsos LLM", "Consistência LLM v1×v2",
             "Erros do auditor v2", "Classificação v2", "Classificação v1", "Gabarito", "Errata"]
    wb._sheets = [wb[n] for n in ordem]
    saida = AQUI / "benchmark_resultados.xlsx"
    wb.save(saida)
    print(f"Planilha salva em {saida}")


if __name__ == "__main__":
    main()
